from io import BytesIO
import re

from docx import Document

from search_anything.embeddings import ContextLimitError
from search_anything.ingestion import (
    TextBlock,
    _docx_blocks,
    _markdown_blocks,
    _python_blocks,
    chunk_blocks,
    extract_document,
)


class WordTokenizer:
    context_tokens = 200

    def count_text_tokens(self, text: str) -> int:
        count = len(text.split()) + 4
        if count > self.context_tokens:
            raise ContextLimitError("too long")
        return count

    def text_token_offsets(self, text: str) -> list[tuple[int, int]]:
        return [(match.start(), match.end()) for match in re.finditer(r"\S+", text)]


def test_markdown_preserves_structural_kinds_and_source_offsets() -> None:
    source = (
        "  # Heading  \n\n"
        "| Name | Value |\n| --- | --- |\n| first | 1 |\n\n"
        "```python\nvalue = 1\n```\n"
    )
    blocks = _markdown_blocks(source)
    table = next(block for block in blocks if block.kind == "table")
    code = next(block for block in blocks if block.kind == "code")
    assert source[table.source_start : table.source_end] == table.text
    assert table.text.startswith("| Name | Value |")
    assert code.text.startswith("```python")


def test_docx_blocks_follow_document_order_and_table_heading() -> None:
    document = Document()
    document.add_heading("Results", level=1)
    document.add_paragraph("Before table")
    document.add_table(rows=1, cols=2).rows[0].cells[0].text = "Name"
    document.tables[0].rows[0].cells[1].text = "Value"
    document.add_paragraph("After table")
    stream = BytesIO()
    document.save(stream)

    blocks = _docx_blocks(stream.getvalue())
    assert [block.kind for block in blocks] == ["heading", "text", "table", "text"]
    assert [block.heading for block in blocks] == ["Results"] * 4
    assert blocks[2].text == "Name | Value"


def test_html_extraction_omits_navigation_and_scripts() -> None:
    html = b"<html><nav>menu</nav><h1>Title</h1><p>Keep this</p><script>secret()</script></html>"
    blocks, image, pages = extract_document("page.html", "text/html", html)
    text = " ".join(block.text for block in blocks)
    assert "Keep this" in text
    assert "menu" not in text
    assert "secret" not in text
    assert image is None
    assert pages is None


def test_python_documents_split_at_function_boundaries() -> None:
    source = (
        "import math\n\ndef first():\n    return math.pi\n\nclass Second:\n    pass\n"
    )
    blocks = _python_blocks(source)
    assert [block.heading for block in blocks] == [
        "Python module imports",
        "Python FunctionDef: first",
        "Python ClassDef: Second",
    ]
    assert (
        source[blocks[1].source_start : blocks[1].source_end].strip() == blocks[1].text
    )


def test_long_table_chunks_repeat_header_and_keep_source_rows() -> None:
    header = "| Name | Value |\n| --- | --- |\n"
    rows = "".join(f"| item-{index} | value-{index} |\n" for index in range(30))
    source = header + rows
    block = TextBlock(
        source,
        "Inventory",
        {"table_id": 4},
        source_start=0,
        source_end=len(source),
        kind="table",
    )
    chunks = chunk_blocks([block], "report.md", WordTokenizer(), 40, 0)
    assert len(chunks) > 1
    assert chunks[0].text.startswith(header)
    for index, chunk in enumerate(chunks[1:], 1):
        assert chunk.text.startswith(header)
        assert chunk.location["repeated_header"] is True
        assert source[chunk.source_start : chunk.source_end].strip() != chunk.text
        assert chunk.token_count <= 40


def test_oversized_prose_ends_at_sentence_boundaries_and_keeps_offsets() -> None:
    source = "A clear sentence ends here. " * 30
    block = TextBlock(
        source,
        None,
        {"offset_unit": "unicode_codepoint"},
        source_start=0,
        source_end=len(source),
    )
    chunks = chunk_blocks([block], None, WordTokenizer(), 18, 0)
    assert len(chunks) > 1
    for chunk in chunks:
        assert source[chunk.source_start : chunk.source_end] == chunk.text
        assert chunk.text.endswith(".")
        assert chunk.token_count <= 18
